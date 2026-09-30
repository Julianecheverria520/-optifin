"""Acceso centralizado al archivo Excel que hace de base de datos local.

Todos los routers leen y escriben a través de este módulo; al migrar a
Supabase basta con reemplazar estas funciones.
"""
import os
import threading
from contextlib import contextmanager
from datetime import date, datetime

import openpyxl
import pandas as pd
from fastapi import HTTPException

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_FILE = os.environ.get("OPTIFIN_EXCEL", os.path.join(BASE_DIR, "OptiFin_Estructura_Datos.xlsx"))

# FastAPI ejecuta los endpoints síncronos en varios hilos: solo uno puede escribir a la vez
_lock = threading.Lock()

MSG_BLOQUEADO = ("No se pudo acceder al archivo Excel: está abierto o bloqueado. "
                 "Ciérralo en Excel o espera a que OneDrive termine de sincronizar.")


# ---------- Lectura ----------

def leer_hoja(hoja: str) -> pd.DataFrame:
    """Lee una hoja completa. Las celdas vacías quedan como NaN, pero textos como
    "N/A" se conservan (por defecto pandas los convierte en NaN)."""
    try:
        df = pd.read_excel(EXCEL_FILE, sheet_name=hoja, keep_default_na=False, na_values=[""])
    except FileNotFoundError:
        raise HTTPException(status_code=500, detail=f"No se encontró el archivo de datos: {EXCEL_FILE}")
    except PermissionError:
        raise HTTPException(status_code=409, detail=MSG_BLOQUEADO)
    return df.dropna(how="all")


def _a_json(valor):
    if pd.isna(valor):
        return None
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    if hasattr(valor, "item"):  # tipos numpy -> tipos de Python
        valor = valor.item()
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def a_registros(df: pd.DataFrame) -> list[dict]:
    """Convierte un DataFrame en una lista de dicts serializable a JSON."""
    return [{k: _a_json(v) for k, v in fila.items()} for fila in df.to_dict(orient="records")]


# ---------- Escritura ----------

def _cargar_libro():
    try:
        return openpyxl.load_workbook(EXCEL_FILE)
    except FileNotFoundError:
        raise HTTPException(status_code=500, detail=f"No se encontró el archivo de datos: {EXCEL_FILE}")
    except PermissionError:
        raise HTTPException(status_code=409, detail=MSG_BLOQUEADO)


@contextmanager
def leer_libro():
    """Abre el Excel solo para consultar (con las mismas funciones que al editar), sin guardar."""
    with _lock:
        yield _cargar_libro()


@contextmanager
def editar_libro():
    """Abre el Excel para escritura (una petición a la vez) y lo guarda al salir.
    Si ocurre un error dentro del bloque, no se guarda nada."""
    with _lock:
        wb = _cargar_libro()
        yield wb
        try:
            wb.save(EXCEL_FILE)
        except PermissionError:
            raise HTTPException(status_code=409, detail=MSG_BLOQUEADO)


def _columnas(ws) -> dict:
    """Mapa {encabezado: número de columna} según la fila 1."""
    return {c.value: c.column for c in ws[1] if c.value}


def _col(ws, nombre: str) -> int:
    cols = _columnas(ws)
    if nombre not in cols:
        raise HTTPException(status_code=500, detail=f"La hoja '{ws.title}' no tiene la columna '{nombre}'")
    return cols[nombre]


def siguiente_id(ws) -> int:
    ids = [v for v in (ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1))
           if isinstance(v, (int, float))]
    return int(max(ids)) + 1 if ids else 1


def buscar_fila(ws, id_registro: int):
    """Número de fila cuyo ID (columna 1) coincide, o None."""
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, (int, float)) and int(v) == int(id_registro):
            return r
    return None


def fila_o_404(ws, id_registro: int, entidad: str) -> int:
    fila = buscar_fila(ws, id_registro)
    if fila is None:
        raise HTTPException(status_code=404, detail=f"{entidad} {id_registro} no encontrado(a)")
    return fila


def valor(ws, fila: int, columna: str):
    return ws.cell(row=fila, column=_col(ws, columna)).value


def agregar(ws, datos: dict) -> int:
    """Agrega una fila poniendo cada valor en la columna con su mismo nombre.
    Asigna y devuelve el nuevo ID (columna 1)."""
    nuevo_id = siguiente_id(ws)
    fila = ws.max_row + 1
    ws.cell(row=fila, column=1).value = nuevo_id
    actualizar(ws, fila, datos)
    return nuevo_id


def actualizar(ws, fila: int, datos: dict):
    for columna, v in datos.items():
        ws.cell(row=fila, column=_col(ws, columna)).value = v


def eliminar(ws, fila: int):
    ws.delete_rows(fila)


def asegurar_estructura(requeridas: dict[str, list[str]]):
    """Crea las hojas y columnas que falten (migración simple del Excel).
    Solo guarda el archivo si hubo cambios."""
    with _lock:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        cambios = []
        for hoja, columnas in requeridas.items():
            if hoja not in wb.sheetnames:
                wb.create_sheet(hoja).append(columnas)
                cambios.append(f"hoja {hoja}")
                continue
            ws = wb[hoja]
            existentes = _columnas(ws)
            for columna in columnas:
                if columna not in existentes:
                    ws.cell(row=1, column=ws.max_column + 1).value = columna
                    cambios.append(f"{hoja}.{columna}")
        if cambios:
            wb.save(EXCEL_FILE)
        return cambios


def existe_valor(ws, columna: str, buscado, id_usuario=None) -> bool:
    """True si alguna fila tiene `buscado` en la columna indicada (solo filas de ese usuario, si se indica)."""
    c = _col(ws, columna)
    c_usuario = _col(ws, "ID_Usuario") if id_usuario is not None else None
    return any(ws.cell(row=r, column=c).value == buscado
               and (c_usuario is None or ws.cell(row=r, column=c_usuario).value == id_usuario)
               for r in range(2, ws.max_row + 1))


# ---------- Datos por usuario ----------

def solo_usuario(df: pd.DataFrame, id_usuario: int) -> pd.DataFrame:
    """Filas de un DataFrame que pertenecen al usuario."""
    return df[df["ID_Usuario"] == id_usuario]


def fila_propia(ws, id_registro, id_usuario: int):
    """Fila del registro si existe y pertenece al usuario; si no, None."""
    fila = buscar_fila(ws, id_registro) if id_registro is not None else None
    if fila is None or valor(ws, fila, "ID_Usuario") != id_usuario:
        return None
    return fila


def fila_propia_o_404(ws, id_registro: int, id_usuario: int, entidad: str) -> int:
    """Como fila_o_404, pero un registro de otro usuario se trata como inexistente."""
    fila = fila_propia(ws, id_registro, id_usuario)
    if fila is None:
        raise HTTPException(status_code=404, detail=f"{entidad} {id_registro} no encontrado(a)")
    return fila
